# -*- coding: utf-8 -*-
"""Contrato HTTP do fluxo independente; nenhum PDF, CSV ou XLSX em disco."""

import sys

sys.dont_write_bytecode = True

import builtins
import io
import json
import re
import tempfile
import unittest
import unicodedata
from contextlib import contextmanager
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / "web"))

import fitz
from flask import Flask, render_template, template_rendered
from openpyxl import load_workbook

import comparador_bp as bp
from extrator_documento_precificado import extrair_documento_precificado
from extrator_orcamento_gamatec import extrair_orcamento_gamatec

UPLOAD = "/comparador/documento-precificado"
PROCESSAR = UPLOAD + "/processar"
RESULTADO = UPLOAD + "/resultado"
DOWNLOAD = UPLOAD + "/download"
FILEID = re.compile(r"sess_precificado_[0-9a-f]{32}\.json\Z")
API = {
    "precificado_upload": (UPLOAD, "GET"),
    "precificado_processar": (PROCESSAR, "POST"),
    "precificado_resultado": (RESULTADO, "GET"),
    "precificado_download": (DOWNLOAD, "GET"),
}
CABECALHO = "COD KRONA descrição VALOR unit qnd TOTAL"
LINHA = "1232 item 1,412 750 R$ 1.059,10"
ORCAMENTO = "1232 ELETRODUTO 1 750 1,500 1.125,00"
TOTAIS = {
    "TOTAL_DOCUMENTO": Decimal("1059.10"),
    "TOTAL_CALCULADO": Decimal("1059.00"),
    "DIFERENCA_TOTAL": Decimal("0.10"),
}


def pdf_em_memoria(*paginas):
    """None gera uma página com imagem, sem camada de texto (necessita OCR)."""
    with fitz.open() as documento:
        for texto in paginas:
            pagina = documento.new_page(width=1200, height=1600)
            if texto is None:
                imagem = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
                imagem.clear_with(220)
                pagina.insert_image(pagina.rect, stream=imagem.tobytes("png"))
            elif texto:
                pagina.insert_text((40, 55), texto, fontsize=10, fontname="helv")
        return documento.tobytes()


def documento_pdf(*linhas):
    return pdf_em_memoria("\n".join((CABECALHO, *(linhas or (LINHA,)))))


def normalizar(texto):
    return "".join(c for c in unicodedata.normalize("NFD", texto.lower())
                   if not unicodedata.combining(c))


class HTMLVisivel(HTMLParser):
    """Ignora scripts/estilos e permite verificar inputs e células reais."""

    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.inputs = []
        self.forms = []
        self.textos = []
        self.headers = []
        self.rows = []
        self._ignorados = 0
        self._row = None
        self._cell = None
        self._cell_tag = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        atributos = dict(attrs)
        if tag in ("script", "style", "template"):
            self._ignorados += 1
        if tag == "input":
            self.inputs.append(atributos)
        if tag == "form":
            self.forms.append(atributos)
        if tag == "tr":
            self._row = []
        if tag in ("td", "th"):
            self._cell = []
            self._cell_tag = tag

    def handle_endtag(self, tag):
        if tag in ("script", "style", "template"):
            self._ignorados = max(0, self._ignorados - 1)
        if tag in ("td", "th") and self._cell is not None:
            texto = " ".join(" ".join(self._cell).split())
            if self._row is not None:
                self._row.append(texto)
            if self._cell_tag == "th":
                self.headers.append(texto)
            self._cell = None
        if tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if not self._ignorados and data.strip():
            self.textos.append(data.strip())
            if self._cell is not None:
                self._cell.append(data.strip())

    @property
    def texto(self):
        return " ".join(self.textos)


@contextmanager
def capturar_templates(app):
    capturados = []

    def receber(sender, template, context, **extra):
        capturados.append((template, context))

    template_rendered.connect(receber, app, weak=False)
    try:
        yield capturados
    finally:
        template_rendered.disconnect(receber, app)


def api_ausente(app):
    regras = {(regra.rule, metodo, regra.endpoint)
              for regra in app.url_map.iter_rules() for metodo in regra.methods}
    return [nome for nome, (rota, metodo) in API.items()
            if not callable(getattr(bp, nome, None))
            or (rota, metodo, "comparador." + nome) not in regras]


class BaseInterface(unittest.TestCase):
    def setUp(self):
        temporario = TemporaryDirectory(prefix="teste_precificado_")
        self.addCleanup(temporario.cleanup)
        self.tmp = Path(temporario.name).resolve()
        self.patch_session = patch.object(bp, "SESSION_TMP_DIR", str(self.tmp))
        self.patch_session.start()
        self.addCleanup(self.patch_session.stop)
        self.legado_save = self._patch(
            bp, "_save_session_data",
            side_effect=AssertionError("Não executar o gravador antigo de sessão."),
        )
        self.legado_csv = self._patch(
            bp, "_csv_oc",
            side_effect=AssertionError("O fluxo independente não pode ler uma OC."),
        )
        self.legado_planilha = self._patch(
            bp, "gerar_planilha_desconto",
            side_effect=AssertionError("Não gravar XLSX/CSV do fluxo antigo."),
        )
        self._patch(bp, "SAIDA_DIR", self.tmp / "ocs_nao_criar")
        self._patch(bp.pd, "read_csv",
                    side_effect=AssertionError("Não ler CSVs vivos ou tabelas."))
        # XMLs transitórios internos do XlsxWriter também ficam no sandbox.
        self._patch(tempfile, "tempdir", str(self.tmp))
        # Defesa adicional: até uma regressão não pode gravar fora do sandbox.
        for modulo in (builtins, io):
            abrir = modulo.open

            def abrir_seguro(file, mode="r", *args, _abrir=abrir, **kwargs):
                if isinstance(file, (str, bytes, Path)) and any(c in mode for c in "wax+"):
                    caminho = Path(file.decode() if isinstance(file, bytes) else file).resolve()
                    if not caminho.is_relative_to(self.tmp):
                        raise AssertionError(f"Escrita fora da sessão temporária: {caminho}")
                return _abrir(file, mode, *args, **kwargs)

            self._patch(modulo, "open", side_effect=abrir_seguro)
        self.app = Flask(
            "teste_interface_precificado",
            template_folder=str(BASE / "web" / "templates"),
            static_folder=str(BASE / "web" / "static"),
        )
        self.app.config.update(TESTING=True, SECRET_KEY="chave-apenas-de-teste")
        self.app.add_url_rule("/", endpoint="login", view_func=lambda: "Login")
        self.app.register_blueprint(bp.comparador_bp)
        self.client = self.app.test_client()
        self.autenticar(self.client)
        self.upload_original = ""

    def _patch(self, alvo, nome, *args, **kwargs):
        remendo = patch.object(alvo, nome, *args, **kwargs)
        objeto = remendo.start()
        self.addCleanup(remendo.stop)
        return objeto

    def autenticar(self, client, user="usuario_teste"):
        with client.session_transaction() as sessao:
            sessao["user"] = user

    def arquivos(self):
        return sorted(p.relative_to(self.tmp).as_posix() for p in self.tmp.rglob("*"))

    def token(self, client=None):
        client = client or self.client
        resposta = client.get(UPLOAD)
        self.assertEqual(resposta.status_code, 200)
        html = HTMLVisivel(resposta.get_data(as_text=True))
        self.upload_original = html.texto
        with client.session_transaction() as sessao:
            token = sessao.get("precificado_csrf")
        self.assertIsInstance(token, str)
        self.assertTrue(token)
        campos = [campo for campo in html.inputs if campo.get("name") == "_csrf"]
        self.assertEqual(len(campos), 1)
        self.assertEqual(campos[0].get("type"), "hidden")
        self.assertEqual(campos[0].get("value"), token)
        return token

    def enviar(self, client=None, documento=None, orcamento=None,
               nome_documento="documento.pdf", nome_orcamento="orcamento.pdf",
               csrf=None, omitir=()):
        client = client or self.client
        dados = {}
        if csrf is not None:
            dados["_csrf"] = csrf
        if "documento" not in omitir:
            dados["documento"] = (io.BytesIO(documento if documento is not None
                                               else documento_pdf()), nome_documento)
        if "orcamento" not in omitir:
            dados["orcamento"] = (io.BytesIO(orcamento if orcamento is not None
                                               else pdf_em_memoria(ORCAMENTO)), nome_orcamento)
        return client.post(PROCESSAR, data=dados, content_type="multipart/form-data")

    def sucesso(self, client=None, **kwargs):
        client = client or self.client
        resposta = self.enviar(client=client, csrf=self.token(client), **kwargs)
        self.assertEqual(resposta.status_code, 303, resposta.get_data(as_text=True))
        destino = urlsplit(resposta.headers["Location"])
        self.assertEqual(destino.path, RESULTADO)
        consulta = parse_qs(destino.query)
        self.assertEqual(set(consulta), {"id"})
        self.assertEqual(len(consulta["id"]), 1)
        fileid = consulta["id"][0]
        self.assertRegex(fileid, FILEID)
        with client.session_transaction() as sessao:
            self.assertEqual(sessao.get("comparador_precificado"), fileid)
        self.assertTrue((self.tmp / fileid).is_file())
        self.assertTrue(all(FILEID.fullmatch(p) for p in self.arquivos()), self.arquivos())
        self.legado_save.assert_not_called()
        self.legado_csv.assert_not_called()
        self.legado_planilha.assert_not_called()
        return fileid

    def erro_upload(self, resposta, templates):
        self.assertEqual(resposta.status_code, 422, resposta.get_data(as_text=True))
        self.assertTrue(templates, "O erro deve renderizar upload, não redirecionar.")
        self.assertIn("upload", templates[-1][0].name.lower())
        html = HTMLVisivel(resposta.get_data(as_text=True))
        nomes = {campo.get("name") for campo in html.inputs}
        self.assertTrue({"documento", "orcamento", "_csrf"}.issubset(nomes))
        self.assertNotEqual(html.texto, self.upload_original, "Erro não visível no upload.")
        self.assertRegex(normalizar(html.texto),
                         r"erro|invalid|selecione|obrigator|nenhum|nao|repetid|duplicad|ocr|layout")

    def negar_id(self, rota, fileid, client=None):
        client = client or self.client
        resposta = client.get(rota, query_string={"id": fileid} if fileid is not None else {})
        self.assertEqual(resposta.status_code, 302)
        self.assertEqual(urlsplit(resposta.headers["Location"]).path, UPLOAD)
        with client.session_transaction() as sessao:
            self.assertTrue(sessao.get("_flashes"), "A recusa deve explicar a sessão inválida.")
        self.assertNotIn(b"item", resposta.data)
        self.assertNotIn(b"PK\x03\x04", resposta.data)
        return resposta

    def planilha(self, fileid, client=None):
        resposta = (client or self.client).get(DOWNLOAD, query_string={"id": fileid})
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.mimetype,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", resposta.headers.get("Content-Disposition", ""))
        self.assertIn(".xlsx", resposta.headers.get("Content-Disposition", ""))
        self.assertTrue(resposta.data.startswith(b"PK\x03\x04"))
        workbook = load_workbook(io.BytesIO(resposta.data), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["comparacao", "gamatec"])
        return workbook

    def registros_excel(self, sheet):
        linhas = list(sheet.iter_rows())
        self.assertTrue(linhas)
        cabecalhos = [cell.value for cell in linhas[0]]
        self.assertEqual(len(cabecalhos), len(set(cabecalhos)))
        return cabecalhos, [dict(zip(cabecalhos, linha)) for linha in linhas[1:]
                           if any(cell.value is not None for cell in linha)]


class TestContratoAPI(BaseInterface):
    def test_funcoes_e_rotas_publicas_exatas(self):
        faltantes = api_ausente(self.app)
        self.assertFalse(faltantes, "API precificada ainda ausente: " + ", ".join(faltantes))
        for nome in API:
            self.assertIs(self.app.view_functions["comparador." + nome], getattr(bp, nome))

    def test_fixtures_reais_sem_orcamento_externo(self):
        documento = extrair_documento_precificado(documento_pdf()).iloc[0]
        self.assertEqual(documento["valor_unitario_oc"], Decimal("1.412"))
        self.assertEqual(documento["quantidade_final"], Decimal("750"))
        self.assertEqual(documento["total_documento"], TOTAIS["TOTAL_DOCUMENTO"])
        itens = extrair_orcamento_gamatec(pdf_em_memoria(ORCAMENTO))["itens"]
        self.assertEqual(len(itens), 1)
        self.assertEqual(itens[0]["codigo"], "1232")
        self.assertEqual(itens[0]["qtde"], 750)
        self.assertEqual(itens[0]["preco_orcamento"], 1.5)
        self.assertEqual(self.arquivos(), [])


class TestFluxoPrecificado(BaseInterface):
    def setUp(self):
        super().setUp()
        faltantes = api_ausente(self.app)
        if faltantes:
            self.skipTest("API ainda ausente: " + ", ".join(faltantes))

    def test_upload_csrf_e_formulario_multipart(self):
        with capturar_templates(self.app) as templates:
            token = self.token()
        self.assertGreaterEqual(len(token), 32)
        self.assertIn("upload", templates[-1][0].name.lower())
        html = HTMLVisivel(self.client.get(UPLOAD).get_data(as_text=True))
        for nome in ("documento", "orcamento"):
            campos = [campo for campo in html.inputs if campo.get("name") == nome]
            self.assertEqual(len(campos), 1)
            self.assertEqual(campos[0].get("type"), "file")
        forms = [form for form in html.forms if form.get("action") == PROCESSAR]
        self.assertEqual(len(forms), 1)
        self.assertEqual(forms[0].get("method", "").lower(), "post")
        self.assertEqual(forms[0].get("enctype"), "multipart/form-data")
        self.assertEqual(self.arquivos(), [])

    def test_todas_rotas_exigem_autenticacao_antes_de_processar(self):
        anonimo = self.app.test_client()
        for rota, metodo in API.values():
            with self.subTest(rota=rota):
                resposta = anonimo.open(rota, method=metodo)
                self.assertEqual(resposta.status_code, 302)
                self.assertEqual(urlsplit(resposta.headers["Location"]).path, "/")
                with anonimo.session_transaction() as sessao:
                    self.assertNotIn("precificado_csrf", sessao)
                    self.assertNotIn("comparador_precificado", sessao)
        self.assertEqual(self.arquivos(), [])

    def test_csrf_ausente_invalido_ou_sem_token_na_sessao(self):
        valido = self.token()
        for token in (None, "token-invalido", valido + "x"):
            with self.subTest(token=token):
                resposta = self.enviar(csrf=token)
                self.assertEqual(resposta.status_code, 400)
                self.assertEqual(self.arquivos(), [])
        with self.client.session_transaction() as sessao:
            sessao.pop("precificado_csrf", None)
        self.assertEqual(self.enviar(csrf=valido).status_code, 400)
        self.assertEqual(self.arquivos(), [])

    def test_csrf_sem_uploads_nao_grava_nem_chama_extratores(self):
        valido = self.token()
        with patch.object(bp, "extrair_orcamento_gamatec") as extrair:
            for token in (None, "errado"):
                with self.subTest(token=token):
                    dados = {} if token is None else {"_csrf": token}
                    resposta = self.client.post(PROCESSAR, data=dados)
                    self.assertEqual(resposta.status_code, 400)
            extrair.assert_not_called()
        with capturar_templates(self.app) as templates:
            resposta = self.client.post(PROCESSAR, data={"_csrf": valido})
        self.erro_upload(resposta, templates)
        self.assertEqual(self.arquivos(), [])
        with self.client.session_transaction() as sessao:
            self.assertNotIn("comparador_precificado", sessao)

    def test_token_de_outra_sessao_nao_e_aceito(self):
        token = self.token()
        outro = self.app.test_client()
        self.autenticar(outro, "outro_usuario")
        self.assertNotEqual(token, self.token(outro))
        self.assertEqual(self.enviar(client=outro, csrf=token).status_code, 400)
        self.assertEqual(self.arquivos(), [])

    def test_campos_ausentes_ou_filename_vazio(self):
        token = self.token()
        casos = ({"omitir": ("documento",)}, {"omitir": ("orcamento",)},
                 {"nome_documento": ""}, {"nome_orcamento": ""})
        for kwargs in casos:
            with self.subTest(caso=kwargs), capturar_templates(self.app) as templates:
                resposta = self.enviar(csrf=token, **kwargs)
                self.erro_upload(resposta, templates)
                self.assertEqual(self.arquivos(), [])

    def test_extensao_final_pdf_obrigatoria_nos_dois_campos(self):
        token = self.token()
        for campo in ("nome_documento", "nome_orcamento"):
            for nome in ("arquivo.txt", "arquivo.pdf.exe", "arquivo", "arquivo.pdf "):
                with self.subTest(campo=campo, nome=nome), capturar_templates(self.app) as templates:
                    resposta = self.enviar(csrf=token, **{campo: nome})
                    self.erro_upload(resposta, templates)
                    self.assertEqual(self.arquivos(), [])

    def test_uuid_seguro_aleatorio_sem_nome_de_upload_ou_usuario(self):
        self.autenticar(self.client, "../user../")
        nomes = {"nome_documento": "../../nomeDocumento.PdF",
                 "nome_orcamento": "nomeOrcamento.PDF"}
        primeiro = self.sucesso(**nomes)
        segundo = self.sucesso(**nomes)
        self.assertNotEqual(primeiro, segundo)
        for fileid in (primeiro, segundo):
            for marcador in ("nomeDocumento", "nomeOrcamento", "user", "..", "/", "\\"):
                self.assertNotIn(marcador, fileid)
        self.assertTrue(all(FILEID.fullmatch(p) for p in self.arquivos()))

    def test_pdfs_invalidos_em_qualquer_campo(self):
        token = self.token()
        for campo in ("documento", "orcamento"):
            for payload in (b"", b"nao e PDF", b"%PDF-1.7\ntruncado"):
                with self.subTest(campo=campo, payload=payload), capturar_templates(self.app) as templates:
                    resposta = self.enviar(csrf=token, **{campo: payload})
                    self.erro_upload(resposta, templates)
                    self.assertEqual(self.arquivos(), [])

    def test_documento_sem_itens_layout_errado_ocr_e_duplicados(self):
        token = self.token()
        validos = CABECALHO + "\n" + LINHA
        casos = {
            "sem_itens": pdf_em_memoria(CABECALHO),
            "layout": pdf_em_memoria("CODIGO DESCRICAO QTD PRECO TOTAL\n" + LINHA),
            "ocr": pdf_em_memoria(None),
            "ocr_apos_item": pdf_em_memoria(validos, None),
            "duplicado": documento_pdf(LINHA, "1232 SEGUNDO 1,00 1 R$ 1,00"),
            "duplicado_normalizado": documento_pdf(LINHA, "0123 A 1,00 1 R$ 1,00",
                                                   "123 B 1,00 1 R$ 1,00"),
            "linha_parcial": documento_pdf(LINHA, "1233 LINHA INCOMPLETA"),
        }
        for nome, payload in casos.items():
            with self.subTest(caso=nome), capturar_templates(self.app) as templates:
                resposta = self.enviar(csrf=token, documento=payload)
                self.erro_upload(resposta, templates)
                self.assertEqual(self.arquivos(), [])
                with self.client.session_transaction() as sessao:
                    self.assertNotIn("comparador_precificado", sessao)

    def test_orcamento_sem_itens_e_codigo_repetido_rejeitados(self):
        token = self.token()
        for texto in ("Total: 0,00", ORCAMENTO + "\n" + ORCAMENTO):
            with self.subTest(texto=texto), capturar_templates(self.app) as templates:
                resposta = self.enviar(csrf=token, orcamento=pdf_em_memoria(texto))
                self.erro_upload(resposta, templates)
                self.assertEqual(self.arquivos(), [])

    def test_falha_de_comparacao_nao_publica_arquivo_parcial(self):
        token = self.token()
        with patch.object(bp, "comparar_documento_precificado",
                          side_effect=ValueError("Erro controlado na comparação.")) as comparar:
            with capturar_templates(self.app) as templates:
                resposta = self.enviar(csrf=token)
            comparar.assert_called_once()
        self.erro_upload(resposta, templates)
        self.assertEqual(self.arquivos(), [])
        with self.client.session_transaction() as sessao:
            self.assertNotIn("comparador_precificado", sessao)

    def test_falha_durante_gravacao_json_remove_arquivo_parcial(self):
        token = self.token()

        def interromper_gravacao(dados, arquivo, *args, **kwargs):
            arquivo.write('{"registros": [')
            raise OSError("Falha de escrita simulada após iniciar o JSON.")

        with patch.object(bp.json, "dump", side_effect=interromper_gravacao) as gravar:
            resposta = self.enviar(csrf=token)
            gravar.assert_called_once()
        self.assertGreaterEqual(resposta.status_code, 400)
        self.assertLess(resposta.status_code, 600)
        with self.client.session_transaction() as sessao:
            self.assertNotIn("comparador_precificado", sessao)
        self.assertEqual(self.arquivos(), [], "JSON parcial não deve ficar publicado no disco.")

    def test_falha_na_publicacao_atomica_preserva_resultado_anterior(self):
        fileid = self.sucesso()
        anteriores = self.arquivos()
        token = self.token()
        with patch.object(bp.os, "replace", side_effect=OSError("Disco indisponível.")):
            resposta = self.enviar(csrf=token)
        self.assertEqual(resposta.status_code, 503)
        self.assertEqual(self.arquivos(), anteriores)
        with self.client.session_transaction() as sessao:
            self.assertEqual(sessao.get("comparador_precificado"), fileid)

    def test_csrf_unicode_invalido_retorna_400(self):
        self.token()
        self.assertEqual(self.enviar(csrf="inválido").status_code, 400)
        self.assertEqual(self.arquivos(), [])

    def test_upload_grande_mantem_formulario_e_resultado_anterior(self):
        fileid = self.sucesso()
        anteriores = self.arquivos()
        token = self.token()
        self.app.config["MAX_CONTENT_LENGTH"] = 100
        with capturar_templates(self.app) as templates:
            resposta = self.enviar(csrf=token)
        self.assertEqual(resposta.status_code, 413)
        self.assertIn("upload", templates[-1][0].name)
        self.assertIn("limite de upload", resposta.get_data(as_text=True))
        self.assertEqual(self.arquivos(), anteriores)
        with self.client.session_transaction() as sessao:
            self.assertEqual(sessao.get("comparador_precificado"), fileid)

    def test_orcamento_numeros_nao_finitos_rejeitados(self):
        token = self.token()
        for campo in ("preco_orcamento", "qtde"):
            for valor in (float("nan"), float("inf")):
                with self.subTest(campo=campo, valor=valor):
                    orcamento = extrair_orcamento_gamatec(pdf_em_memoria(ORCAMENTO))
                    orcamento["itens"][0][campo] = valor
                    with patch.object(bp, "extrair_orcamento_gamatec", return_value=orcamento):
                        with capturar_templates(self.app) as templates:
                            resposta = self.enviar(csrf=token)
                    self.erro_upload(resposta, templates)
                    self.assertEqual(self.arquivos(), [])

    def test_erro_novo_nao_substitui_resultado_existente_por_parcial(self):
        fileid = self.sucesso()
        anteriores = self.arquivos()
        token = self.token()
        with capturar_templates(self.app) as templates:
            resposta = self.enviar(csrf=token, documento=documento_pdf(LINHA, "1233 INCOMPLETO"))
        self.erro_upload(resposta, templates)
        self.assertEqual(self.arquivos(), anteriores)
        with self.client.session_transaction() as sessao:
            self.assertEqual(sessao.get("comparador_precificado"), fileid)

    def test_json_decimal_string_nan_none_e_totais_documentais(self):
        fileid = self.sucesso(orcamento=pdf_em_memoria(
            ORCAMENTO + "\n1233 OUTRO ITEM 1 1 2,000 2,00"))
        with (self.tmp / fileid).open(encoding="utf-8") as arquivo:
            texto = arquivo.read()

        def rejeitar_constante(valor):
            raise AssertionError("Constante JSON não finita: " + valor)

        dados = json.loads(texto, parse_constant=rejeitar_constante)
        registros = {r["CODIGO"]: r for r in dados["registros"]}
        for coluna, esperado in TOTAIS.items():
            valor = registros["1232"][coluna]
            self.assertIsInstance(valor, str)
            self.assertEqual(Decimal(valor), esperado)
            self.assertIsNone(registros["1233"][coluna])
        preco_documento = ("PRECO_DOCUMENTO" if "PRECO_DOCUMENTO" in registros["1233"]
                           else "PRECO_OC")
        self.assertIsNone(registros["1233"][preco_documento])
        self.assertIsNone(registros["1233"]["DESCONTO"])

    def test_resultado_contexto_titulos_precos_e_totais_originais(self):
        fileid = self.sucesso()
        with capturar_templates(self.app) as templates:
            resposta = self.client.get(RESULTADO, query_string={"id": fileid})
        self.assertEqual(resposta.status_code, 200)
        self.assertTrue(templates)
        contexto = templates[-1][1]
        self.assertIs(contexto.get("precificado"), True)
        registro = next(r for r in contexto["registros"] if r["CODIGO"] == "1232")
        for coluna, esperado in TOTAIS.items():
            self.assertEqual(Decimal(str(registro[coluna])), esperado)
        html = HTMLVisivel(resposta.get_data(as_text=True))
        headers = [normalizar(h) for h in html.headers]
        for palavras in (("preco", "documento"), ("total", "documento"),
                         ("total", "calculado"), ("diferenca",)):
            self.assertTrue(any(all(p in h for p in palavras) for h in headers), headers)
        self.assertNotIn("Preço OC", html.headers)
        self.assertNotIn("Total OC", html.headers)
        linha = next(row for row in html.rows if row and row[0] == "1232")
        texto = " ".join(linha)
        self.assertRegex(texto, r"1[.,]412\b")
        self.assertRegex(texto, r"1[.,]500\b")
        self.assertRegex(texto, r"1[.,]?059[.,]10\b")
        self.assertRegex(texto, r"1[.,]?059[.,]00\b")
        self.assertRegex(texto, r"0[.,]10\b")

    def test_ids_ausentes_invalidos_e_traversal_inclusive_iguais_a_sessao(self):
        valido = self.sucesso()
        ids = (None, "", "sess_precificado_" + "a" * 32 + ".json",
               "../outro.json", "..\\outro.json", str(self.tmp / "valido.json"),
               str(self.tmp / valido), "../" + self.tmp.name + "/" + valido,
               ".\\" + valido,
               "sess_precificado_" + "A" * 32 + ".json",
               "sess_precificado_" + "a" * 31 + ".json",
               "sess_precificado_" + "a" * 32 + ".json.bak")
        for rota in (RESULTADO, DOWNLOAD):
            for fileid in ids:
                with self.subTest(rota=rota, fileid=fileid):
                    with self.client.session_transaction() as sessao:
                        sessao["comparador_precificado"] = valido
                    self.negar_id(rota, fileid)
                    # Mesmo igual ao cookie, traversal não pode acessar o JSON existente.
                    with self.client.session_transaction() as sessao:
                        sessao["comparador_precificado"] = fileid
                    self.negar_id(rota, fileid)

    def test_sessao_expirada_com_usuario_presente(self):
        fileid = self.sucesso()
        with self.client.session_transaction() as sessao:
            sessao.pop("comparador_precificado")
        for rota in (RESULTADO, DOWNLOAD):
            with self.subTest(rota=rota):
                self.negar_id(rota, fileid)

    def test_arquivo_ausente_mantendo_id_valido_na_sessao(self):
        fileid = self.sucesso()
        (self.tmp / fileid).unlink()
        for rota in (RESULTADO, DOWNLOAD):
            with self.subTest(rota=rota):
                self.negar_id(rota, fileid)
        self.assertEqual(self.arquivos(), [])

    def test_id_de_outra_sessao_nao_le_arquivo_existente(self):
        meu_id = self.sucesso()
        outro = self.app.test_client()
        self.autenticar(outro, "outro_usuario")
        outro_id = self.sucesso(client=outro)
        self.assertNotEqual(meu_id, outro_id)
        abrir_real = builtins.open
        io_abrir_real = io.open
        acessados = []

        def vigiar(abrir):
            def abrir_vigiado(file, *args, **kwargs):
                if isinstance(file, (str, Path)):
                    acessados.append(Path(file).resolve())
                return abrir(file, *args, **kwargs)
            return abrir_vigiado

        with patch.object(builtins, "open", side_effect=vigiar(abrir_real)), \
                patch.object(io, "open", side_effect=vigiar(io_abrir_real)):
            for rota in (RESULTADO, DOWNLOAD):
                with self.subTest(rota=rota):
                    self.negar_id(rota, outro_id)
                    self.negar_id(rota, meu_id, client=outro)
        self.assertNotIn(self.tmp / meu_id, acessados)
        self.assertNotIn(self.tmp / outro_id, acessados)
        self.assertEqual(self.client.get(RESULTADO, query_string={"id": meu_id}).status_code, 200)
        self.assertEqual(outro.get(RESULTADO, query_string={"id": outro_id}).status_code, 200)

    def test_download_dois_sheets_numeros_totais_e_desconto_string(self):
        fileid = self.sucesso()
        antes = self.arquivos()
        wb = self.planilha(fileid)
        headers, audit = self.registros_excel(wb["comparacao"])
        self.assertIn("PRECO_DOCUMENTO", headers)
        self.assertNotIn("PRECO_OC", headers)
        self.assertEqual(len(audit), 1)
        registro = audit[0]
        for coluna, esperado in {"PRECO_DOCUMENTO": Decimal("1.412"), **TOTAIS}.items():
            cell = registro[coluna]
            self.assertEqual(cell.data_type, "n")
            self.assertIsInstance(cell.value, (int, float))
            self.assertEqual(Decimal(str(cell.value)), esperado)
        for coluna in headers:
            if str(coluna).startswith("PRECO") and registro[coluna].value is not None:
                self.assertIsInstance(registro[coluna].value, (int, float))
        headers, gamatec = self.registros_excel(wb["gamatec"])
        self.assertEqual(headers, ["DESCRICAO", "CODIGO", "QUANTIDADE", "DESCONTO"])
        self.assertEqual(len(gamatec), 1)
        self.assertEqual(str(gamatec[0]["CODIGO"].value), "1232")
        self.assertEqual(gamatec[0]["QUANTIDADE"].value, 750)
        desconto = gamatec[0]["DESCONTO"]
        self.assertEqual(desconto.data_type, "s")
        self.assertEqual(desconto.value, "5.86667")
        self.assertEqual(self.arquivos(), antes, "Download não deve criar XLSX/CSV em disco.")

    def test_codigo_0647_preco_e_desconto_corretos_no_resultado_e_excel(self):
        fileid = self.sucesso(
            documento=documento_pdf("0647 LUVA ESGOTO PRIM. - DN 100 2,226 1.000 R$ 2.226,00"),
            orcamento=pdf_em_memoria(
                "Produto Emb. Qtde Unitário Total IPI ICMS Ret Unit.Imp. Tot.c/Imp.\n"
                "0647 LUVA ESGOTO PRIM. - DN 100 10 1.000 2,738 2.738,00 0,00 0,00 2,738 2.738,00"
            ),
        )
        dados = json.loads((self.tmp / fileid).read_text(encoding="utf-8"))
        registro = dados["registros"][0]
        self.assertEqual(registro["CODIGO"], "647")
        self.assertEqual(registro["QUANTIDADE"], 1000)
        self.assertEqual(registro["PRECO_ORCAMENTO"], 2.738)
        self.assertEqual(registro["DESCONTO"], 18.69979)
        resposta = self.client.get(RESULTADO, query_string={"id": fileid})
        self.assertEqual(resposta.status_code, 200)
        html = HTMLVisivel(resposta.get_data(as_text=True))
        self.assertIn("2.738", html.texto)
        self.assertIn("18.69979", html.texto)
        wb = self.planilha(fileid)
        _, auditoria = self.registros_excel(wb["comparacao"])
        self.assertEqual(auditoria[0]["PRECO_ORCAMENTO"].value, 2.738)
        self.assertEqual(auditoria[0]["QUANTIDADE"].value, 1000)
        self.assertEqual(auditoria[0]["DESCONTO"].value, 18.69979)
        _, gamatec = self.registros_excel(wb["gamatec"])
        self.assertEqual(gamatec[0]["DESCONTO"].value, "18.69979")

    def test_orcamento_coluna_invalida_mostra_erro_e_preserva_resultado_anterior(self):
        fileid = self.sucesso()
        antes = self.arquivos()
        with capturar_templates(self.app) as templates:
            resposta = self.enviar(
                csrf=self.token(),
                orcamento=pdf_em_memoria(
                    "Produto Emb. Qtde Unitário Total IPI ICMS Ret Unit.Imp. Tot.c/Imp.\n"
                    "1232 PRIMEIRO ITEM 1 750 1,500 1.125,00 0,00 0,00 1,500 1.125,00\n"
                    "0647 LUVA DN 100 10 1.000 INVALIDO 2.738,00 0,00 0,00 2,738 2.738,00"
                ),
            )
        self.erro_upload(resposta, templates)
        self.assertIn("Colunas inválidas no item 647", resposta.get_data(as_text=True))
        self.assertEqual(self.arquivos(), antes)
        with self.client.session_transaction() as sessao:
            self.assertEqual(sessao["comparador_precificado"], fileid)

    def test_download_audita_todos_e_filtra_somente_ok_abaixo(self):
        fileid = self.sucesso(
            documento=documento_pdf(LINHA, "1233 ABAIXO 2,000 1 R$ 2,00",
                                    "1234 SEM ORCAMENTO 1,000 1 R$ 1,00"),
            orcamento=pdf_em_memoria(ORCAMENTO + "\n1233 ABAIXO 1 1 1,000 1,00"
                                    + "\n1235 SOMENTE ORCAMENTO 1 1 2,000 2,00"),
        )
        wb = self.planilha(fileid)
        _, audit = self.registros_excel(wb["comparacao"])
        estados = {str(r["CODIGO"].value): r["STATUS"].value for r in audit}
        self.assertEqual(estados, {"1232": "OK", "1233": "ABAIXO",
                                  "1234": "SEM_ORC", "1235": "SEM_OC"})
        _, gamatec = self.registros_excel(wb["gamatec"])
        descontos = {str(r["CODIGO"].value): r["DESCONTO"].value for r in gamatec}
        self.assertEqual(descontos, {"1232": "5.86667", "1233": "0.00000"})

    def test_download_sem_correspondencia_audit_completa_gamatec_vazio(self):
        fileid = self.sucesso(orcamento=pdf_em_memoria("1233 OUTRO 1 1 2,000 2,00"))
        wb = self.planilha(fileid)
        _, audit = self.registros_excel(wb["comparacao"])
        self.assertEqual({str(r["CODIGO"].value) for r in audit}, {"1232", "1233"})
        self.assertEqual({r["STATUS"].value for r in audit}, {"SEM_ORC", "SEM_OC"})
        documento = next(r for r in audit if str(r["CODIGO"].value) == "1232")
        for coluna, esperado in TOTAIS.items():
            self.assertEqual(Decimal(str(documento[coluna].value)), esperado)
        headers, gamatec = self.registros_excel(wb["gamatec"])
        self.assertEqual(headers, ["DESCRICAO", "CODIGO", "QUANTIDADE", "DESCONTO"])
        self.assertEqual(gamatec, [])

    def test_formulas_e_urls_externas_permanecem_texto_sem_hyperlink(self):
        payloads = {"1232": "=2+2", "1233": "https://example.invalid/externo"}
        pdf = documento_pdf(*(f"{codigo} {descricao} 1,412 750 R$ 1.059,10"
                              for codigo, descricao in payloads.items()))
        extraido = extrair_documento_precificado(pdf)
        self.assertEqual(extraido["descricao_krona"].tolist(), list(payloads.values()))
        fileid = self.sucesso(documento=pdf, orcamento=pdf_em_memoria(
            ORCAMENTO + "\n1233 ELETRODUTO 1 750 1,500 1.125,00"))
        wb = self.planilha(fileid)
        for sheet in wb:
            with self.subTest(sheet=sheet.title):
                _, registros = self.registros_excel(sheet)
                self.assertEqual(len(registros), 2)
                for registro in registros:
                    codigo = str(registro["CODIGO"].value)
                    descricao = registro["DESCRICAO"]
                    self.assertEqual(descricao.value, payloads[codigo])
                    self.assertEqual(descricao.data_type, "s")
                for linha in sheet:
                    for cell in linha:
                        self.assertNotEqual(cell.data_type, "f", cell.coordinate)
                        self.assertIsNone(cell.hyperlink, cell.coordinate)


class TestRegressaoInterfaceLegada(BaseInterface):
    def test_oc_inexistente_mantem_redirect_legado(self):
        self.legado_csv.side_effect = None
        self.legado_csv.return_value = None
        resposta = self.client.get("/comparador/inexistente")
        self.assertEqual(resposta.status_code, 302)
        self.assertEqual(urlsplit(resposta.headers["Location"]).path, "/")
        self.legado_csv.assert_called_once_with("inexistente")
        self.assertEqual(self.arquivos(), [])
        self.legado_save.assert_not_called()

    def test_render_resultado_sem_precificado_mantem_preco_oc_total_oc(self):
        registro = {
            "CODIGO": "1232", "DESCRICAO": "ITEM LEGADO", "QUANTIDADE": 750,
            "PRECO_OC": 1.412, "PRECO_ORCAMENTO": 1.5, "DESCONTO": 5.86667,
            "PRECO_FINAL": "1.41", "STATUS": "OK", "OBS": None,
        }
        dados = {"numero_orcamento": "TESTE", "emissao": "19/10/2026"}
        stats = {"total": 1, "ok": 1, "abaixo": 0, "sem_orc": 0,
                 "sem_oc": 0, "para_gamatec": 1}
        with self.app.test_request_context("/comparador/OC_TESTE/resultado"):
            html = render_template("resultado.html", oc_id="OC_TESTE", dados=dados,
                                   registros=[registro], stats=stats)
        pagina = HTMLVisivel(html)
        self.assertIn("Preço OC", pagina.headers)
        self.assertIn("Total OC", pagina.headers)
        self.assertNotIn("Preço Documento", pagina.headers)
        self.assertRegex(pagina.texto, r"1[.,]412\b")
        self.assertRegex(pagina.texto, r"1[.,]?059[.,]00\b")
        self.assertEqual(self.arquivos(), [])
        self.legado_save.assert_not_called()


if __name__ == "__main__":
    unittest.main(verbosity=2)
