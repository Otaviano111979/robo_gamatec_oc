"""Testa os estados visuais da planilha de digitacao manual."""

import os
import sys
import tempfile

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from processar_oc_individual_pasta_ajustado import gerar_planilha_enxuta_individual


def _item(**kw):
    base = {
        "match_encontrado": False,
        "tipo_match": "",
        "codigo_krona": None,
        "descricao_krona": None,
        "descricao_oc": "",
        "quantidade_final": 1,
        "valor_unitario": None,
        "preco_unit": None,
        "revisao_manual": False,
        "sugestao_composicao": None,
        "candidatos_revisao": None,
        "unidade_venda_krona": None,
    }
    base.update(kw)
    return base


df = pd.DataFrame([
    _item(match_encontrado=True, tipo_match="MATCH_REGRA_CATALOGO", codigo_krona="521",
          descricao_krona="SOLUCAO PREPARADORA 1000 ML", descricao_oc="SOLUCAO LIMPADORA",
          quantidade_final=5, unidade_venda_krona="PC"),
    _item(match_encontrado=True, tipo_match="MATCH_SEMANTICO_REVISAR", codigo_krona="787",
          descricao_krona="TORNEIRA PARA PIA", descricao_oc="SOLUCAO LIMPADORA",
          revisao_manual=True, unidade_venda_krona="PC"),
    _item(descricao_oc="PASTA LUBRIFICANTE", tipo_match="SEM_MATCH", quantidade_final=5),
    _item(tipo_match="SUGESTAO_COMPOSICAO", descricao_oc="TE PVC SOLDAVEL REDUCAO 60X50MM",
          sugestao_composicao="459 TE SOLD - 60MM + 364 BUCHA RED.SOLD.CURTA 60 X 50MM",
          revisao_manual=True, quantidade_final=30),
    _item(match_encontrado=True, tipo_match="MATCH_CODIGO_MRV", codigo_krona="929",
          descricao_krona="TUBO MRV", descricao_oc="TUBO MRV", quantidade_final=2),
    _item(match_encontrado=True, tipo_match="MATCH_DESCRICAO", codigo_krona="23",
          descricao_krona="TUBO PVC SOLDAVEL - 6M - 20MM", descricao_oc="PLUG 3/4",
          quantidade_final=200, valor_unitario=1.43, unidade_venda_krona="BR"),
    _item(match_encontrado=True, tipo_match="MATCH_DESCRICAO_REVISAR", codigo_krona="1288",
          descricao_krona="QUADRO DISTR 36D/27N SEM BARRAMENTO",
          descricao_oc="QUADRO 36 DIN", revisao_manual=True, unidade_venda_krona="PC"),
])

with tempfile.TemporaryDirectory() as tmp:
    csv = os.path.join(tmp, "planilha.csv")
    xlsx = os.path.join(tmp, "planilha.xlsx")
    out = gerar_planilha_enxuta_individual(df, csv, xlsx)

    assert list(out.columns) == ["DESCRICAO", "CODIGO", "QUANTIDADE"], out.columns
    lido = pd.read_csv(csv, sep=";", encoding="utf-8-sig")
    assert list(lido.columns) == ["DESCRICAO", "CODIGO", "QUANTIDADE"], lido.columns
    assert lido.loc[3, "CODIGO"] == "459 + 364", lido.loc[3, "CODIGO"]
    assert "Revisar" not in lido["DESCRICAO"].str.cat(sep=" ")

    wb = load_workbook(xlsx)
    ws = wb.active
    cores = [ws.cell(row=i, column=1).fill.start_color.rgb for i in range(2, 9)]
    textos = [str(ws.cell(row=i, column=1).value) for i in range(2, 9)]
    codigos = [ws.cell(row=i, column=2).value for i in range(2, 9)]

    assert not cores[0].endswith("FFFF99"), cores[0]
    assert cores[1].endswith("DDEBF7"), cores[1]
    assert cores[2].endswith("FFFF99"), cores[2]
    assert cores[3].endswith("DDEBF7"), cores[3]
    assert cores[4].endswith("FFB347"), cores[4]
    assert not cores[5].endswith("FFB347"), cores[5]
    assert cores[6].endswith("DDEBF7"), cores[6]

    assert "Revisar" in textos[1]
    assert "Revisar" not in textos[2]
    assert "Revisar" in textos[3]
    assert "Verificar preço" in textos[4] and "Revisar" not in textos[4]
    assert "Vendido em BR" in textos[5]
    assert "Revisar" in textos[6]

    assert codigos[3] == "459 + 364", codigos[3]

print("OK: planilha distingue match, revisao, sem match e composicao")
